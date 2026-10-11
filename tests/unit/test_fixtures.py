"""Synthetic fixtures: schemas, built files, eval references, separation, and data review."""

import ast
import datetime as dt
import json
import re
import subprocess
import sys
from pathlib import Path
from typing import Any, Literal

import docx
import openpyxl
import pytest
import yaml
from pydantic import BaseModel, ConfigDict, Field

from stewardry.platform.access import load_access_mapping

ROOT = Path(__file__).resolve().parents[2]
FIXTURES = ROOT / "fixtures"
EVALS = ROOT / "evals"
DATASET = yaml.safe_load((EVALS / "datasets" / "questionnaire-v1.yaml").read_text())
AS_OF = dt.date.fromisoformat(DATASET["as_of"])
PASSAGE = re.compile(r"^(\d+)\.(\d+) (.+)$")


class _Strict(BaseModel):
    model_config = ConfigDict(extra="forbid")


class PolicyFront(_Strict):
    policy_id: str = Field(pattern=r"^SEC-\d{2}$")
    title: str
    version: int = Field(ge=1)
    status: Literal["current", "superseded"]
    classification: Literal["internal", "confidential"]
    owner: str
    effective_date: dt.date
    review_date: dt.date
    confluence_page_id: str = Field(pattern=r"^\d+$")
    superseded_by: str | None = None


def policies() -> dict[str, tuple[PolicyFront, dict[str, str], str]]:
    found = {}
    for path in sorted((FIXTURES / "policies").glob("SEC-*.md")):
        _, front, body = path.read_text().split("---", 2)
        meta = PolicyFront.model_validate(yaml.safe_load(front))
        passages = {}
        for line in body.splitlines():
            match = PASSAGE.match(line)
            if match:
                passages[f"{match.group(1)}.{match.group(2)}"] = match.group(3)
        found[meta.policy_id] = (meta, passages, body)
    return found


POLICIES = policies()


def questionnaires() -> dict[str, dict[str, Any]]:
    out = {}
    for path in sorted((FIXTURES / "questionnaires").glob("QN-*.yaml")):
        data: dict[str, Any] = yaml.safe_load(path.read_text())
        out[data["questionnaire_id"]] = data
    return out


QUESTIONNAIRES = questionnaires()


# ------------------------------------------------------------------- policies


def test_there_are_21_well_formed_policies() -> None:
    assert len(POLICIES) == 21
    for policy_id, (meta, passages, body) in POLICIES.items():
        number = int(policy_id.split("-")[1])
        assert meta.confluence_page_id == str(100 + number)
        assert list(passages) == [f"{number}.{i}" for i in range(1, len(passages) + 1)]
        assert len(passages) >= 2
        assert meta.review_date > meta.effective_date
        assert "<" not in body and "{" not in body, f"{policy_id}: plain content only, no markup"


def test_supersession_and_staleness_are_as_designed() -> None:
    superseded = {pid for pid, (m, _, _) in POLICIES.items() if m.status == "superseded"}
    assert superseded == {"SEC-21"}
    assert POLICIES["SEC-21"][0].superseded_by == "SEC-05"
    assert POLICIES["SEC-05"][0].status == "current"
    stale = {
        pid
        for pid, (m, _, _) in POLICIES.items()
        if m.status == "current" and m.review_date < AS_OF
    }
    assert stale == {"SEC-19"}
    confidential = {
        pid for pid, (m, _, _) in POLICIES.items() if m.classification == "confidential"
    }
    assert confidential == {"SEC-04", "SEC-09", "SEC-17"}


def test_the_backup_conflict_is_real() -> None:
    assert "35 days" in POLICIES["SEC-07"][1]["7.2"]
    assert "90 days" in POLICIES["SEC-15"][1]["15.3"]


def test_catalog_matches_the_policies() -> None:
    catalog = yaml.safe_load((ROOT / "policies/catalogs/secpol.yaml").read_text())
    by_id = {p["policy_id"]: p for p in catalog["pages"]}
    assert set(by_id) == set(POLICIES)
    for policy_id, (meta, _, _) in POLICIES.items():
        page = by_id[policy_id]
        assert (page["page_id"], page["title"], page["status"], page["classification"]) == (
            meta.confluence_page_id,
            meta.title,
            meta.status,
            meta.classification,
        )


# -------------------------------------------------------------- questionnaires


def test_questionnaires_and_jira_requests_line_up() -> None:
    customers = {
        c.customer_id
        for c in load_access_mapping(ROOT / "deploy/access/access-mapping.example.yaml").customers
    }
    assert set(QUESTIONNAIRES) == {"QN-LARKSPUR-01", "QN-LARKSPUR-02", "QN-QUILLON-01"}
    for data in QUESTIONNAIRES.values():
        assert [q["id"] for q in data["questions"]] == [f"Q{i:02d}" for i in range(1, 31)]
        assert data["customer_id"] in customers
        assert (FIXTURES / "templates" / f"{data['template']}.yaml").exists()
        issue = json.loads((FIXTURES / "jira" / f"{data['jira_issue_key']}.json").read_text())
        assert issue["fields"]["customfield_10042"] == data["customer_id"]


def _xlsx_content(path: Path) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path)
    ws = wb["Questionnaire"]
    return {
        "sheets": wb.sheetnames,
        "cells": [[c.value for c in row] for row in ws.iter_rows()],
        "hidden_columns": sorted(k for k, d in ws.column_dimensions.items() if d.hidden),
        "hidden_rows": sorted(r for r, d in ws.row_dimensions.items() if d.hidden),
        "comments": [
            (c.coordinate, c.comment.text) for row in ws.iter_rows() for c in row if c.comment
        ],
        "formulas": [wb["Summary"]["B2"].value, wb["Summary"]["B3"].value],
    }


def _docx_content(path: Path) -> dict[str, Any]:
    document = docx.Document(str(path))
    table = document.tables[0]
    return {
        "rows": [[cell.text for cell in row.cells] for row in table.rows],
        "hidden": [
            r.text
            for row in table.rows
            for c in row.cells
            for p in c.paragraphs
            for r in p.runs
            if r.font.hidden
        ],
        "comments": [c.text for c in document.comments],
        "paragraphs": [p.text for p in document.paragraphs],
    }


def test_committed_office_files_match_a_fresh_build(tmp_path: Path) -> None:
    subprocess.run(  # noqa: S603 - fixed arguments: this repo's own build script
        [sys.executable, str(ROOT / "tools/build_fixtures.py"), str(tmp_path)],
        check=True,
        capture_output=True,
    )
    for name in ("QN-LARKSPUR-01.xlsx", "QN-LARKSPUR-02.xlsx"):
        assert _xlsx_content(FIXTURES / "questionnaires" / name) == _xlsx_content(tmp_path / name)
    name = "QN-QUILLON-01.docx"
    assert _docx_content(FIXTURES / "questionnaires" / name) == _docx_content(tmp_path / name)


def test_hidden_adversarial_content_is_where_the_yaml_says() -> None:
    first = _xlsx_content(FIXTURES / "questionnaires/QN-LARKSPUR-01.xlsx")
    assert first["hidden_columns"] == ["F"] and len(first["comments"]) == 1
    assert first["formulas"] == ["=COUNTA(Questionnaire!D2:D500)", "=COUNTA(Questionnaire!A2:A500)"]
    second = _xlsx_content(FIXTURES / "questionnaires/QN-LARKSPUR-02.xlsx")
    assert len(second["hidden_rows"]) == 1
    quillon = _docx_content(FIXTURES / "questionnaires/QN-QUILLON-01.docx")
    assert len(quillon["hidden"]) == 1 and len(quillon["comments"]) == 1
    assert len(quillon["rows"]) == 31 and quillon["paragraphs"][-1].startswith("Thank you")


# ------------------------------------------------------------------- evals


def test_every_question_has_exactly_one_eval_item() -> None:
    items = DATASET["items"]
    asked = {(qid, q["id"]) for qid, data in QUESTIONNAIRES.items() for q in data["questions"]}
    graded = [
        (i["questionnaire_id"], i["question_id"])
        for i in items
        if not i["question_id"].startswith("H")
    ]
    assert sorted(graded) == sorted(asked)
    hidden = sum(len(data["hidden"]) for data in QUESTIONNAIRES.values())
    assert sum(1 for i in items if i["question_id"].startswith("H")) == hidden


def _passage(ref: str) -> str:
    policy_id, number = ref.split("#")
    return POLICIES[policy_id][1][number]


def test_every_citation_points_at_a_real_passage() -> None:
    for item in DATASET["items"]:
        for ref in item.get("citations", []) + item.get("must_not_cite", []):
            assert _passage(ref), ref


def test_expected_outcomes_follow_the_sets() -> None:
    expected = {
        "golden": "answer",
        "abstention": "needs_expert",
        "conflict_staleness": "needs_expert",
    }
    for item in DATASET["items"]:
        if item["set"] in expected:
            assert item["expected_outcome"] == expected[item["set"]], item
        if item["set"] == "golden":
            assert item["citations"] and item["key_points"], item
            for ref in item["citations"]:
                meta = POLICIES[ref.split("#")[0]][0]
                assert meta.status == "current" and meta.review_date >= AS_OF, ref
        if item["set"] == "abstention":
            assert "citations" not in item, item
        for ref in item.get("must_not_cite", []):
            assert POLICIES[ref.split("#")[0]][0].status == "superseded", ref


# -------------------------------------------------------------- separation


def _string_literals(root: Path) -> set[str]:
    found: set[str] = set()
    for module in root.rglob("*.py"):
        for node in ast.walk(ast.parse(module.read_text())):
            if isinstance(node, ast.Constant) and isinstance(node.value, str):
                found.add(node.value)
    return found


def test_runtime_code_never_refers_to_the_eval_directory() -> None:
    offenders = {
        s for s in _string_literals(ROOT / "src") if re.search(r"(^|[/\\])evals([/\\]|$)", s)
    }
    assert offenders == set()


@pytest.mark.parametrize("folder", ["fixtures", "policies"])
def test_agent_visible_files_hold_no_expected_answers(folder: str) -> None:
    for path in (ROOT / folder).rglob("*"):
        if path.suffix in {".yaml", ".json", ".md"}:
            text = path.read_text()
            for key in ("expected_outcome", "key_points", "must_not_cite"):
                assert key not in text, (path, key)


# ------------------------------------------------------------- data review


SYNTHETIC_CONTENT = [
    "fixtures/policies",
    "fixtures/questionnaires",
    "fixtures/jira",
    "fixtures/templates",
    "policies/catalogs",
    "evals/datasets",
]


@pytest.mark.parametrize("folder", SYNTHETIC_CONTENT)
def test_data_review_no_branding_emails_or_real_links(folder: str) -> None:
    """Synthetic business content only; config may name the model (technical attribution)."""
    paths = [p for p in (ROOT / folder).rglob("*") if p.suffix in {".yaml", ".json", ".md"}]
    assert paths, folder
    for path in paths:
        text = path.read_text()
        assert not re.search(r"anthropic|claude", text, re.IGNORECASE), path
        for email in re.findall(r"[\w.+-]+@([\w-]+\.[\w.-]+)", text):
            assert email.startswith("example."), (path, email)
        for host in re.findall(r"https?://([^/\s\"']+)", text):
            assert host.endswith(("example.com", "example.org")), (path, host)
