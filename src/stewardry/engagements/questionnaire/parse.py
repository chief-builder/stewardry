"""Reading uploaded questionnaires in the supported templates (spec section 5.2).

Uploaded files are untrusted. Before any Office library touches one, the zip
container is checked: size, entry count, total uncompressed size and
compression ratio (zip bombs), no macros, no external links, and the expected
main content type. XML is parsed with entity resolution disabled (openpyxl
through defusedxml; python-docx with resolve_entities=False).

The template decides what a question is. Text anywhere else is never a
question; hidden or secondary text becomes a screening finding. Layouts the
template does not describe are rejected with an explanation the uploader can
act on.
"""

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePath
from typing import Any

import docx
import openpyxl
import openpyxl.xml
from openpyxl.utils import column_index_from_string

from stewardry.engagements.questionnaire.screening import (
    FindingKind,
    ScreeningFinding,
    location_finding,
    screen_text,
)
from stewardry.engagements.questionnaire.templates import docx_template, xlsx_template

MAX_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED = 50 * 1024 * 1024
MAX_ENTRIES = 500
MAX_RATIO = 100
MAX_QUESTIONS = 200
QUESTION_ID = re.compile(r"^[A-Za-z0-9._-]{1,32}$")
MAIN_TYPES = {
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
}
SUPPORTED = "Use the published .xlsx (sqt-xlsx-v1) or .docx (sqt-docx-v1) template."
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class UnsupportedQuestionnaire(Exception):
    """The file cannot be read as a supported questionnaire. The message says why."""


@dataclass(frozen=True)
class Question:
    question_id: str
    ordinal: int
    section: str
    text: str
    location: dict[str, Any]


@dataclass(frozen=True)
class ParsedQuestionnaire:
    template_id: str
    format: str
    sha256: str
    questions: list[Question]
    findings: list[ScreeningFinding] = field(default_factory=list)


# ------------------------------------------------------------- container


def _check_container(filename: str, data: bytes) -> str:
    suffix = PurePath(filename).suffix.lower()
    if suffix in {".xlsm", ".docm", ".xls", ".doc"}:
        raise UnsupportedQuestionnaire(
            f"{suffix} files are not supported (macro-enabled or legacy format). {SUPPORTED}"
        )
    if suffix not in MAIN_TYPES:
        raise UnsupportedQuestionnaire(f"Only .xlsx and .docx files are supported. {SUPPORTED}")
    if not data:
        raise UnsupportedQuestionnaire("The file is empty.")
    if len(data) > MAX_BYTES:
        raise UnsupportedQuestionnaire(f"The file is larger than {MAX_BYTES // 1024 // 1024} MB.")
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise UnsupportedQuestionnaire(
            f"The file is not a valid {suffix} document. {SUPPORTED}"
        ) from exc
    with archive:
        entries = archive.infolist()
        if len(entries) > MAX_ENTRIES:
            raise UnsupportedQuestionnaire("The file has too many internal parts.")
        total = sum(e.file_size for e in entries)
        if total > MAX_UNCOMPRESSED or any(
            e.compress_size and e.file_size / e.compress_size > MAX_RATIO for e in entries
        ):
            raise UnsupportedQuestionnaire("The file expands to an unsafe size.")
        names = {e.filename for e in entries}
        if any(n.lower().endswith("vbaproject.bin") for n in names):
            raise UnsupportedQuestionnaire(f"Files with macros are not supported. {SUPPORTED}")
        if any(n.startswith(("xl/externalLinks/", "word/externalLinks/")) for n in names):
            raise UnsupportedQuestionnaire("Files with external links are not supported.")
        try:
            content_types = archive.read("[Content_Types].xml").decode("utf-8", "replace")
        except KeyError as exc:
            raise UnsupportedQuestionnaire(f"The file is not a valid {suffix} document.") from exc
    if MAIN_TYPES[suffix] not in content_types:
        raise UnsupportedQuestionnaire(
            f"The file's contents do not match its {suffix} extension. {SUPPORTED}"
        )
    return suffix


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _check_ids(questions: list[Question]) -> None:
    if not questions:
        raise UnsupportedQuestionnaire("No questions were found in the template's question area.")
    if len(questions) > MAX_QUESTIONS:
        raise UnsupportedQuestionnaire(f"At most {MAX_QUESTIONS} questions are supported.")
    seen: set[str] = set()
    for q in questions:
        if q.question_id in seen:
            raise UnsupportedQuestionnaire(f"Question ID {q.question_id} is used more than once.")
        seen.add(q.question_id)


# ------------------------------------------------------------------ xlsx


def _hidden_columns(ws: Any) -> set[int]:
    hidden: set[int] = set()
    for key, dim in ws.column_dimensions.items():
        if dim.hidden:
            low = dim.min or column_index_from_string(key)
            high = dim.max or low
            hidden.update(range(low, high + 1))
    return hidden


def _parse_xlsx(data: bytes) -> tuple[str, list[Question], list[ScreeningFinding]]:
    template = xlsx_template()
    if not openpyxl.xml.DEFUSEDXML:  # pragma: no cover - guarded by the dependency
        raise RuntimeError("refusing to parse untrusted spreadsheets without defusedxml")
    try:
        wb = openpyxl.load_workbook(
            io.BytesIO(data), read_only=False, data_only=False, keep_links=False
        )
    except Exception as exc:
        raise UnsupportedQuestionnaire(f"The spreadsheet could not be read. {SUPPORTED}") from exc
    sheet_name = template.question_sheet
    if sheet_name not in wb.sheetnames:
        raise UnsupportedQuestionnaire(f"No sheet named '{sheet_name}' was found. {SUPPORTED}")
    ws = wb[sheet_name]
    if ws.sheet_state != "visible":
        raise UnsupportedQuestionnaire(f"The '{sheet_name}' sheet must be visible.")
    cols = template.columns
    template_columns = [cols.id, cols.section, cols.question, cols.answer, cols.notes]
    headers = [_text(ws[f"{c}{template.header_row}"].value) for c in template_columns]
    if headers != template.headers:
        raise UnsupportedQuestionnaire(
            f"Row {template.header_row} of '{sheet_name}' must read: "
            + ", ".join(template.headers)
            + "."
        )

    findings: list[ScreeningFinding] = []
    questions: list[Question] = []
    last_template_col = max(column_index_from_string(c) for c in template_columns)
    row_question: dict[int, str] = {}

    for row in range(template.first_question_row, ws.max_row + 1):
        cells = {c: ws[f"{c}{row}"] for c in (cols.id, cols.section, cols.question)}
        values = {c: _text(cell.value) for c, cell in cells.items()}
        row_text = " ".join(_text(c.value) for c in ws[row] if c.value is not None)
        if ws.row_dimensions[row].hidden:
            if row_text:
                findings.append(location_finding("hidden_row", f"{sheet_name}!row {row}", row_text))
            continue
        if any(cell.data_type == "f" or values[c].startswith("=") for c, cell in cells.items()):
            raise UnsupportedQuestionnaire(
                f"Row {row}: the ID, Section, and Question cells must hold plain text, "
                "not formulas."
            )
        if not values[cols.id] and not values[cols.question]:
            continue
        if not values[cols.id] or not values[cols.question]:
            raise UnsupportedQuestionnaire(f"Row {row} needs both an ID and a question.")
        if not QUESTION_ID.fullmatch(values[cols.id]):
            raise UnsupportedQuestionnaire(
                f"Row {row}: question IDs may use letters, digits, '.', '_' and '-' (up to 32)."
            )
        qid = values[cols.id]
        row_question[row] = qid
        questions.append(
            Question(
                qid,
                len(questions) + 1,
                values[cols.section],
                values[cols.question],
                {
                    "sheet": sheet_name,
                    "row": row,
                    "question_cell": f"{cols.question}{row}",
                    "answer_cell": f"{cols.answer}{row}",
                },
            )
        )
        findings += screen_text(values[cols.question], f"{sheet_name}!{cols.question}{row}", qid)

    for sheet in wb.worksheets:
        hidden_sheet = sheet.sheet_state != "visible"
        sheet_hidden_cols = _hidden_columns(sheet)
        for line in sheet.iter_rows():
            for cell in line:
                text = _text(cell.value)
                ref = f"{sheet.title}!{cell.coordinate}"
                row_number = cell.row or 0
                cell_qid: str | None = row_question.get(row_number) if sheet is ws else None
                if cell.comment is not None and cell.comment.text.strip():
                    findings.append(location_finding("comment", ref, cell.comment.text, cell_qid))
                if not text or (sheet is ws and sheet.row_dimensions[row_number].hidden):
                    continue
                col = cell.column if isinstance(cell.column, int) else 0
                if hidden_sheet or col in sheet_hidden_cols:
                    kind: FindingKind = (
                        "hidden_column" if col in sheet_hidden_cols else "hidden_cell"
                    )
                    findings.append(location_finding(kind, ref, text, cell_qid))
                elif (
                    sheet is ws
                    and row_number >= template.first_question_row
                    and (col > last_template_col)
                ):
                    findings.append(location_finding("outside_template", ref, text, cell_qid))
                elif sheet is not ws:
                    findings += screen_text(text, ref)
    _check_ids(questions)
    return template.template_id, questions, findings


# ------------------------------------------------------------------ docx


def _visible_text(paragraphs: Any) -> str:
    return "\n".join(
        "".join(run.text for run in p.runs if not run.font.hidden) for p in paragraphs
    ).strip()


def _hidden_text(paragraphs: Any) -> str:
    return " ".join(run.text for p in paragraphs for run in p.runs if run.font.hidden).strip()


def _parse_docx(data: bytes) -> tuple[str, list[Question], list[ScreeningFinding]]:
    template = docx_template()
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise UnsupportedQuestionnaire(f"The document could not be read. {SUPPORTED}") from exc
    if len(document.tables) <= template.table_index:
        raise UnsupportedQuestionnaire(f"No question table was found. {SUPPORTED}")
    table = document.tables[template.table_index]
    rows = list(table.rows)
    width = len(template.headers)
    for index, row in enumerate(rows):
        cells = row.cells
        if len(cells) != width or len({id(c._tc) for c in cells}) != width:
            raise UnsupportedQuestionnaire(
                f"Table row {index + 1} must have exactly {width} separate cells "
                "(merged cells are not supported)."
            )
        if any(c.tables for c in cells):
            raise UnsupportedQuestionnaire("Tables inside table cells are not supported.")
    headers = [_text(c.text) for c in rows[template.header_row].cells]
    if headers != template.headers:
        raise UnsupportedQuestionnaire(
            "The first row of the question table must read: " + ", ".join(template.headers) + "."
        )

    cols = template.columns
    findings: list[ScreeningFinding] = []
    questions: list[Question] = []
    for index in range(template.header_row + 1, len(rows)):
        cells = rows[index].cells
        qid = _visible_text(cells[cols.id].paragraphs)
        text = _visible_text(cells[cols.question].paragraphs)
        for cell, name in ((c, n) for n, c in enumerate(cells)):
            hidden = _hidden_text(cell.paragraphs)
            if hidden:
                findings.append(
                    location_finding(
                        "hidden_text",
                        f"table {template.table_index + 1}, row {index + 1}, cell {name + 1}",
                        hidden,
                        qid or None,
                    )
                )
        if not qid and not text:
            continue
        if not qid or not text:
            raise UnsupportedQuestionnaire(
                f"Table row {index + 1} needs both an ID and a question."
            )
        if not QUESTION_ID.fullmatch(qid):
            raise UnsupportedQuestionnaire(
                f"Table row {index + 1}: question IDs may use letters, digits, '.', '_' and '-'."
            )
        questions.append(
            Question(
                qid,
                len(questions) + 1,
                _visible_text(cells[cols.section].paragraphs),
                text,
                {
                    "table": template.table_index,
                    "row": index,
                    "question_cell": [index, cols.question],
                    "answer_cell": [index, cols.answer],
                },
            )
        )
        findings += screen_text(text, f"table row {index + 1}", qid)

    for number, paragraph in enumerate(document.paragraphs, start=1):
        visible = _visible_text([paragraph])
        hidden = _hidden_text([paragraph])
        if hidden:
            findings.append(location_finding("hidden_text", f"paragraph {number}", hidden))
        if visible:
            findings += screen_text(visible, f"paragraph {number}")
    for number, comment in enumerate(document.comments, start=1):
        if comment.text.strip():
            findings.append(location_finding("comment", f"comment {number}", comment.text))
    for section in document.sections:
        for part in (section.header, section.footer):
            for number, paragraph in enumerate(part.paragraphs, start=1):
                visible = _visible_text([paragraph])
                if visible:
                    findings += screen_text(visible, f"header or footer paragraph {number}")
    body = document.element.body
    for tag in ("ins", "del"):
        for element in body.iter(f"{W}{tag}"):
            text = "".join(element.itertext()).strip()
            if text:
                findings.append(location_finding("tracked_change", f"tracked {tag}", text))
    _check_ids(questions)
    return template.template_id, questions, findings


# ------------------------------------------------------------------ entry


def parse_questionnaire(filename: str, data: bytes) -> ParsedQuestionnaire:
    suffix = _check_container(filename, data)
    if suffix == ".xlsx":
        template_id, questions, findings = _parse_xlsx(data)
    else:
        template_id, questions, findings = _parse_docx(data)
    return ParsedQuestionnaire(
        template_id, suffix[1:], hashlib.sha256(data).hexdigest(), questions, findings
    )
