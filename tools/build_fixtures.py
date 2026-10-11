"""Build the synthetic questionnaire files from their YAML sources.

fixtures/questionnaires/<id>.yaml + the package's template definitions
  -> fixtures/questionnaires/<id>.xlsx or .docx

Adversarial content from each file's `hidden:` list is placed where the
supported templates can hide it: a hidden column, a cell comment, a hidden
row, hidden text, or a document comment. Screening (S9) must find all of it.

Usage: python tools/build_fixtures.py [output_dir]
"""

import datetime as dt
import sys
from pathlib import Path
from typing import Any

import docx
import yaml
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
SOURCES = ROOT / "fixtures" / "questionnaires"
TEMPLATES = ROOT / "src" / "stewardry" / "engagements" / "questionnaire" / "templates"
FIXED_TIME = dt.datetime(2026, 10, 1, 9, 0, 0)  # stable document metadata
AUTHOR = "Synthetic customer reviewer"


def _load(path: Path) -> dict[str, Any]:
    data: dict[str, Any] = yaml.safe_load(path.read_text(encoding="utf-8"))
    return data


def build_xlsx(source: dict[str, Any], template: dict[str, Any], out: Path) -> None:
    columns = template["columns"]
    wb = Workbook()
    ws = wb.active
    assert ws is not None  # noqa: S101
    ws.title = template["question_sheet"]
    header_fill = PatternFill("solid", fgColor="DDE4EE")
    for index, title in enumerate(template["headers"], start=1):
        cell = ws.cell(row=template["header_row"], column=index, value=title)
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for letter, width in zip("ABCDE", (8, 18, 70, 60, 30), strict=True):
        ws.column_dimensions[letter].width = width

    hidden = source.get("hidden", [])
    rows: dict[str, int] = {}
    row = template["first_question_row"]
    for question in source["questions"]:
        ws[f"{columns['id']}{row}"] = question["id"]
        ws[f"{columns['section']}{row}"] = question["section"]
        ws[f"{columns['question']}{row}"] = question["text"]
        rows[question["id"]] = row
        row += 1
        for item in hidden:
            if item["kind"] == "hidden_row" and item["after"] == question["id"]:
                ws[f"{columns['question']}{row}"] = item["text"]
                ws.row_dimensions[row].hidden = True
                row += 1

    for item in hidden:
        if item["kind"] == "hidden_column_cell":
            ws[f"F{rows[item['row_of']]}"] = item["text"]
            ws.column_dimensions["F"].hidden = True
        elif item["kind"] == "cell_comment":
            ws[f"{columns['question']}{rows[item['on']]}"].comment = Comment(item["text"], AUTHOR)

    instructions = wb.create_sheet("Instructions")
    instructions["A1"] = source["title"]
    instructions["A2"] = "Answer every question in the Answer column. Do not change other columns."
    instructions["A3"] = "Return the completed file to your account contact."
    summary = wb.create_sheet("Summary")
    summary["A1"], summary["B1"] = "Measure", "Value"
    summary["A2"], summary["A3"] = "Answered", "Questions"
    for target, formula in template["other_sheets"][1]["formulas"].items():
        summary[target] = formula

    wb.properties.creator = AUTHOR
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    wb.save(out)


def build_docx(source: dict[str, Any], template: dict[str, Any], out: Path) -> None:
    document = docx.Document()
    document.add_heading(source["title"], level=1)
    document.add_paragraph("Please answer each question in the Answer column of the table below.")
    headers = template["headers"]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, title in enumerate(headers):
        run = table.rows[0].cells[index].paragraphs[0].add_run(title)
        run.bold = True

    hidden = {item["on"]: item for item in source.get("hidden", [])}
    for question in source["questions"]:
        cells = table.add_row().cells
        cells[template["columns"]["id"]].text = question["id"]
        cells[template["columns"]["section"]].text = question["section"]
        paragraph = cells[template["columns"]["question"]].paragraphs[0]
        run = paragraph.add_run(question["text"])
        item = hidden.get(question["id"])
        if item is not None and item["kind"] == "hidden_text":
            secret = paragraph.add_run(" " + item["text"])
            secret.font.hidden = True
            secret.font.size = Pt(1)
        elif item is not None and item["kind"] == "comment":
            document.add_comment(run, text=item["text"], author=AUTHOR, initials="SC")
    document.add_paragraph("Thank you. Return this document to your account contact.")

    document.core_properties.author = AUTHOR
    document.core_properties.created = FIXED_TIME
    document.core_properties.modified = FIXED_TIME
    document.core_properties.last_modified_by = AUTHOR
    document.save(str(out))


def build_all(output_dir: Path) -> list[Path]:
    built = []
    for source_path in sorted(SOURCES.glob("QN-*.yaml")):
        source = _load(source_path)
        template = _load(TEMPLATES / f"{source['template']}.yaml")
        out = output_dir / f"{source['questionnaire_id']}.{template['format']}"
        if template["format"] == "xlsx":
            build_xlsx(source, template, out)
        else:
            build_docx(source, template, out)
        built.append(out)
    return built


def main() -> int:
    output_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else SOURCES
    for path in build_all(output_dir):
        print(f"built {path.relative_to(ROOT) if path.is_relative_to(ROOT) else path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
